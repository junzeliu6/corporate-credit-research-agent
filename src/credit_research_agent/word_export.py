"""Convert an existing Markdown research report into a business-style Word file.

This module is deliberately local and deterministic: it does not invoke an LLM,
perform network access, or rewrite the report's analysis.  It only translates
supported Markdown structures and inserts already-generated local images.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
import re
import tempfile
from typing import Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


_IMAGE_PATTERN = re.compile(r"^\s*!\[(?P<alt>[^]]*)]\((?P<target>[^)]+)\)\s*$")
_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_UL_PATTERN = re.compile(r"^\s*[-+*]\s+(.+)$")
_OL_PATTERN = re.compile(r"^\s*(?P<number>\d+)[.)]\s+(?P<text>.+)$")
_TABLE_SEPARATOR_CELL = re.compile(r"^:?-{3,}:?$")
_INLINE_PATTERN = re.compile(
    r"(`[^`]+`|\*\*[^*]+\*\*|__[^_]+__|(?<![\w*])\*[^*]+\*(?![\w*])|(?<![\w_])_[^_]+_(?![\w_])|\[[^]]+]\([^)]+\))"
)


@dataclass(frozen=True)
class WordExportResult:
    markdown_path: Path
    docx_path: Path
    image_count: int
    table_count: int
    heading_count: int
    warning_codes: tuple[str, ...] = ()


@dataclass(frozen=True)
class MarkdownImageInspection:
    """Validated local image dependencies referenced by a Markdown report."""

    image_paths: tuple[Path, ...]
    requires_ownership_image: bool
    ownership_image_path: Path | None


class WordExportError(RuntimeError):
    """A safe, user-facing Word conversion error."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass
class _ExportStats:
    image_count: int = 0
    table_count: int = 0
    heading_count: int = 0
    ownership_image_inserted: bool = False
    has_ownership_heading: bool = False
    warnings: list[str] | None = None

    def __post_init__(self) -> None:
        if self.warnings is None:
            self.warnings = []


def export_markdown_report(
    markdown_path: str | Path,
    output_path: str | Path | None = None,
    *,
    allowed_root: str | Path | None = None,
    ownership_image_path: str | Path | None = None,
    company_name: str | None = None,
    ticker: str | None = None,
    report_date: str | date | None = None,
    require_ownership_image: bool = False,
) -> WordExportResult:
    """Export an existing Markdown report as DOCX without changing its prose.

    ``allowed_root`` is mandatory for UI callers and constrains the Markdown,
    output, and every inserted image to a confirmed research-task directory.
    CLI callers may omit it when deliberately exporting a local report.
    """

    source = Path(markdown_path).expanduser().resolve()
    if not source.is_file():
        raise WordExportError("markdown_not_found", f"Markdown 报告不存在：{source}")
    if source.suffix.lower() not in {".md", ".markdown"}:
        raise WordExportError("unsupported_markdown_extension", "Word 导出仅接受 Markdown 文件。")

    root = Path(allowed_root).expanduser().resolve() if allowed_root is not None else source.parent
    _ensure_within(source, root, code="markdown_outside_allowed_root")

    destination = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else source.with_suffix(".docx").resolve()
    )
    if destination.suffix.lower() != ".docx":
        raise WordExportError("invalid_output_extension", "Word 输出文件必须使用 .docx 扩展名。")
    _ensure_within(destination, root, code="output_outside_allowed_root")

    explicit_image = None
    if ownership_image_path is not None:
        explicit_image = Path(ownership_image_path).expanduser().resolve()
        _ensure_within(explicit_image, root, code="image_outside_allowed_root")
        if not explicit_image.is_file():
            raise WordExportError("ownership_image_missing", f"股权结构图不存在：{explicit_image}")

    try:
        markdown = source.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise WordExportError("markdown_encoding_error", "Markdown 报告不是有效的 UTF-8 文本。") from exc
    if not markdown.strip():
        raise WordExportError("empty_markdown", "Markdown 报告为空，无法导出正式 Word 报告。")

    document = Document()
    _configure_document(document)
    title = _first_heading(markdown) or "企业授信研究报告"
    _add_cover(
        document,
        title=title,
        company_name=company_name,
        ticker=ticker,
        report_date=report_date,
    )
    document.add_page_break()

    stats = _ExportStats()
    _render_markdown(
        document,
        markdown,
        source_dir=source.parent,
        allowed_root=root,
        explicit_ownership_image=explicit_image,
        stats=stats,
        skip_first_title=True,
    )
    if explicit_image is not None and not stats.ownership_image_inserted:
        _add_figure(document, explicit_image, "股权结构图")
        stats.image_count += 1
        stats.ownership_image_inserted = True
        stats.warnings.append("ownership_image_appended")
    if require_ownership_image and not stats.ownership_image_inserted:
        raise WordExportError(
            "ownership_image_required",
            "报告包含股权研究内容，但未找到可嵌入的股权结构图。请检查任务目录中的 PNG 文件。",
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent, delete=False
        ) as handle:
            temp_path = Path(handle.name)
        document.save(temp_path)
        # Re-open before delivery so corrupt packages fail inside the exporter.
        Document(temp_path)
        temp_path.replace(destination)
    except WordExportError:
        raise
    except Exception as exc:  # python-docx exposes multiple low-level exception types
        raise WordExportError("docx_write_failed", f"Word 文件生成失败：{exc}") from exc
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)

    return WordExportResult(
        markdown_path=source,
        docx_path=destination,
        image_count=stats.image_count,
        table_count=stats.table_count,
        heading_count=stats.heading_count,
        warning_codes=tuple(stats.warnings),
    )


def inspect_markdown_images(
    markdown_path: str | Path,
    *,
    allowed_root: str | Path | None = None,
    ownership_image_path: str | Path | None = None,
) -> MarkdownImageInspection:
    """Resolve every local image used by the exporter without writing a DOCX.

    UI callers use this before returning a cached Word file so changes or
    deletion of any referenced chart invalidate the cache.  The same path
    rules as :func:`export_markdown_report` apply.
    """

    source = Path(markdown_path).expanduser().resolve()
    if not source.is_file():
        raise WordExportError("markdown_not_found", f"Markdown 报告不存在：{source}")
    root = Path(allowed_root).expanduser().resolve() if allowed_root is not None else source.parent
    _ensure_within(source, root, code="markdown_outside_allowed_root")
    explicit_image = None
    if ownership_image_path is not None:
        explicit_image = Path(ownership_image_path).expanduser().resolve()
        _ensure_within(explicit_image, root, code="image_outside_allowed_root")
        if not explicit_image.is_file():
            raise WordExportError("ownership_image_missing", f"股权结构图不存在：{explicit_image}")
    try:
        markdown = source.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise WordExportError("markdown_encoding_error", "Markdown 报告不是有效的 UTF-8 文本。") from exc

    dependencies: list[Path] = []
    requires_ownership = explicit_image is not None
    resolved_ownership_image: Path | None = None
    for line in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        heading = _HEADING_PATTERN.match(line)
        if heading and "股权结构图" in _strip_inline_markdown(heading.group(2)):
            requires_ownership = True
        image = _IMAGE_PATTERN.match(line)
        if image is None:
            continue
        target = _parse_image_target(image.group("target"))
        is_ownership = _is_ownership_image(target, image.group("alt") or "")
        requires_ownership = requires_ownership or is_ownership
        dependency = (
            explicit_image
            if explicit_image is not None and is_ownership
            else _resolve_image(target, source_dir=source.parent, allowed_root=root)
        )
        if is_ownership:
            resolved_ownership_image = dependency
        if dependency not in dependencies:
            dependencies.append(dependency)
    if explicit_image is not None and explicit_image not in dependencies:
        dependencies.append(explicit_image)
    if explicit_image is not None:
        resolved_ownership_image = explicit_image
    return MarkdownImageInspection(
        image_paths=tuple(dependencies),
        requires_ownership_image=requires_ownership,
        ownership_image_path=resolved_ownership_image,
    )


def _ensure_within(path: Path, root: Path, *, code: str) -> None:
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise WordExportError(code, f"路径不在允许的研究任务目录内：{path}") from exc


def _configure_document(document: DocxDocument) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    _set_east_asia_font(normal, "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    title = styles["Title"]
    title.font.name = "黑体"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    _set_east_asia_font(title, "黑体")
    title_border = title.element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        title.element.pPr.remove(title_border)

    for index, size in ((1, 16), (2, 14), (3, 12), (4, 11)):
        style = styles[f"Heading {index}"]
        style.font.name = "黑体"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        _set_east_asia_font(style, "黑体")
        style.paragraph_format.space_before = Pt(12 if index == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "宋体"
        style.font.size = Pt(10.5)
        _set_east_asia_font(style, "宋体")
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.left_indent = Cm(0.8)
        style.paragraph_format.space_after = Pt(2)

    if "Report Caption" not in styles:
        caption = styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.base_style = styles["Normal"]
        caption.font.name = "宋体"
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = RGBColor(89, 89, 89)
        _set_east_asia_font(caption, "宋体")
        caption.paragraph_format.first_line_indent = None
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "企业授信研究报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.font.name = "宋体"
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor(127, 127, 127)
    _set_run_east_asia_font(header_run, "宋体")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    _set_run_east_asia_font(run, "宋体")
    _append_field(run, "PAGE")
    run = footer.add_run(" 页")
    _set_run_east_asia_font(run, "宋体")


def _add_cover(
    document: DocxDocument,
    *,
    title: str,
    company_name: str | None,
    ticker: str | None,
    report_date: str | date | None,
) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(88)
    paragraph.paragraph_format.space_after = Pt(32)
    paragraph.add_run(_strip_inline_markdown(title))

    metadata = [
        ("公司名称", company_name or _company_from_title(title) or "未注明"),
        ("证券代码", ticker or "未注明"),
        ("研究日期", _format_report_date(report_date)),
    ]
    for label, value in metadata:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(8)
        label_run = p.add_run(f"{label}：")
        label_run.bold = True
        _set_run_east_asia_font(label_run, "黑体")
        value_run = p.add_run(value)
        _set_run_east_asia_font(value_run, "宋体")

    statement = document.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.first_line_indent = None
    statement.paragraph_format.space_before = Pt(46)
    statement.add_run("本报告由 AI 辅助生成，相关结论需人工复核。")


def _render_markdown(
    document: DocxDocument,
    markdown: str,
    *,
    source_dir: Path,
    allowed_root: Path,
    explicit_ownership_image: Path | None,
    stats: _ExportStats,
    skip_first_title: bool,
) -> None:
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    first_heading_skipped = False
    in_fence = False
    fence_lines: list[str] = []
    has_ownership_image_reference = any(
        match is not None
        and _is_ownership_image(
            _parse_image_target(match.group("target")), match.group("alt") or ""
        )
        for line in lines
        for match in [_IMAGE_PATTERN.match(line)]
    )

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_fence:
                _add_code_block(document, fence_lines)
                fence_lines.clear()
                in_fence = False
            else:
                in_fence = True
            index += 1
            continue
        if in_fence:
            fence_lines.append(raw)
            index += 1
            continue
        if not stripped or stripped == "---" or stripped == "***":
            index += 1
            continue

        heading = _HEADING_PATTERN.match(raw)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if skip_first_title and not first_heading_skipped and level == 1:
                first_heading_skipped = True
                index += 1
                continue
            paragraph = document.add_paragraph(style=f"Heading {min(level, 4)}")
            paragraph.paragraph_format.first_line_indent = None
            _add_inline_runs(paragraph, text)
            stats.heading_count += 1
            clean_heading = _strip_inline_markdown(text)
            if "股权结构" in clean_heading:
                stats.has_ownership_heading = True
                if (
                    explicit_ownership_image is not None
                    and not has_ownership_image_reference
                    and not stats.ownership_image_inserted
                    and "股权结构图" in clean_heading
                ):
                    _add_figure(document, explicit_ownership_image, "股权结构图")
                    stats.image_count += 1
                    stats.ownership_image_inserted = True
            index += 1
            continue

        image = _IMAGE_PATTERN.match(raw)
        if image:
            target = _parse_image_target(image.group("target"))
            is_ownership_image = _is_ownership_image(target, image.group("alt") or "")
            # A task manager may supply the already validated PNG explicitly.
            # It replaces the Markdown ownership image at the same document
            # position; other local charts continue to follow Markdown paths.
            image_path = (
                explicit_ownership_image
                if explicit_ownership_image is not None and is_ownership_image
                else _resolve_image(target, source_dir=source_dir, allowed_root=allowed_root)
            )
            _add_figure(document, image_path, image.group("alt") or "研究图表")
            stats.image_count += 1
            if is_ownership_image:
                stats.ownership_image_inserted = True
            index += 1
            continue

        if _is_table_start(lines, index):
            header = _parse_table_row(lines[index])
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and _looks_like_table_row(lines[index]):
                rows.append(_parse_table_row(lines[index]))
                index += 1
            _add_table(document, header, rows)
            stats.table_count += 1
            continue

        unordered = _UL_PATTERN.match(raw)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, unordered.group(1))
            index += 1
            continue

        ordered = _OL_PATTERN.match(raw)
        if ordered:
            start_at = int(ordered.group("number"))
            items: list[str] = []
            while index < len(lines):
                item = _OL_PATTERN.match(lines[index])
                if item is None:
                    break
                items.append(item.group("text"))
                index += 1
            num_id = _create_numbering_instance(document, start_at=start_at)
            for item_text in items:
                paragraph = document.add_paragraph(style="List Number")
                _apply_numbering(paragraph, num_id)
                _add_inline_runs(paragraph, item_text)
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_after = Pt(4)
            _add_inline_runs(paragraph, stripped.lstrip("> "))
            _set_paragraph_left_rule(paragraph)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            value = candidate.strip()
            if (
                not value
                or value.startswith("```")
                or _HEADING_PATTERN.match(candidate)
                or _IMAGE_PATTERN.match(candidate)
                or _UL_PATTERN.match(candidate)
                or _OL_PATTERN.match(candidate)
                or value.startswith(">")
                or value in {"---", "***"}
                or _is_table_start(lines, index)
            ):
                break
            paragraph_lines.append(value)
            index += 1
        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, _join_soft_lines(paragraph_lines))

    if in_fence and fence_lines:
        _add_code_block(document, fence_lines)


def _first_heading(markdown: str) -> str | None:
    for line in markdown.splitlines():
        match = _HEADING_PATTERN.match(line)
        if match and len(match.group(1)) == 1:
            return _strip_inline_markdown(match.group(2))
    return None


def _company_from_title(title: str) -> str | None:
    clean = _strip_inline_markdown(title)
    for suffix in ("企业授信研究报告", "授信研究报告", "研究报告"):
        if clean.endswith(suffix):
            candidate = clean[: -len(suffix)].strip(" —-：:")
            return candidate or None
    return None


def _format_report_date(value: str | date | None) -> str:
    if value is None:
        return date.today().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() or date.today().isoformat()


def _parse_image_target(raw: str) -> str:
    target = raw.strip()
    if target.startswith("<") and ">" in target:
        target = target[1 : target.index(">")]
    elif " \"" in target:
        target = target.split(" \"", 1)[0]
    return target.strip()


def _resolve_image(target: str, *, source_dir: Path, allowed_root: Path) -> Path:
    lowered = target.lower()
    if lowered.startswith(("http://", "https://", "data:", "file://")):
        raise WordExportError("remote_image_not_allowed", f"Word 导出不读取远程图片：{target}")
    candidate = (source_dir / target).resolve()
    _ensure_within(candidate, allowed_root, code="image_outside_allowed_root")
    if candidate.suffix.lower() == ".svg":
        png_candidate = candidate.with_suffix(".png")
        if png_candidate.is_file():
            candidate = png_candidate
    if not candidate.is_file():
        raise WordExportError("referenced_image_missing", f"Markdown 引用的图片不存在：{candidate}")
    if candidate.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}:
        raise WordExportError("unsupported_image_type", f"Word 导出暂不支持该图片格式：{candidate.suffix}")
    return candidate


def _is_ownership_image(target: str, alt: str) -> bool:
    return "ownership" in Path(target).name.lower() or "股权" in alt


def _add_figure(document: DocxDocument, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.keep_with_next = True
    try:
        image = DocxImage.from_file(str(image_path))
        max_width = Cm(15.4)
        max_height = Cm(20.0)
        scaled_height = image.height * max_width / image.width
        if scaled_height > max_height:
            paragraph.add_run().add_picture(str(image_path), height=max_height)
        else:
            paragraph.add_run().add_picture(str(image_path), width=max_width)
    except Exception as exc:
        raise WordExportError("image_embed_failed", f"图片无法嵌入 Word：{image_path.name}（{exc}）") from exc
    if caption.strip():
        cap = document.add_paragraph(style="Report Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(_strip_inline_markdown(caption.strip()))


def _looks_like_table_row(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _parse_table_row(line: str) -> list[str]:
    value = line.strip().strip("|")
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", value)]


def _is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or not _looks_like_table_row(lines[index]):
        return False
    cells = _parse_table_row(lines[index + 1])
    return bool(cells) and all(_TABLE_SEPARATOR_CELL.match(cell.replace(" ", "")) for cell in cells)


def _add_table(document: DocxDocument, header: list[str], rows: list[list[str]]) -> None:
    column_count = max(len(header), *(len(row) for row in rows)) if rows else len(header)
    table = document.add_table(rows=1, cols=column_count)
    table.alignment = 1
    table.autofit = True
    _set_table_borders(table)
    _set_repeat_table_header(table.rows[0])

    for column, value in enumerate(header):
        cell = table.rows[0].cells[column]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "1F4E78")
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        run = p.add_run(_strip_inline_markdown(value))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
        _set_run_east_asia_font(run, "黑体")

    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column in range(column_count):
            value = values[column] if column < len(values) else ""
            cell = cells[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            if row_index % 2 == 0:
                _set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if _looks_numeric(value) else WD_ALIGN_PARAGRAPH.LEFT
            _add_inline_runs(p, value, default_size=9)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_inline_runs(paragraph, text: str, *, default_size: float | None = None) -> None:
    cursor = 0
    for match in _INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            _append_styled_run(paragraph, text[cursor : match.start()], size=default_size)
        token = match.group(0)
        if token.startswith(("**", "__")):
            run = _append_styled_run(paragraph, token[2:-2], size=default_size)
            run.bold = True
        elif token.startswith("`"):
            run = _append_styled_run(paragraph, token[1:-1], size=default_size, font="Consolas")
            run.font.color.rgb = RGBColor(89, 89, 89)
        elif token.startswith(("*", "_")):
            run = _append_styled_run(paragraph, token[1:-1], size=default_size)
            run.italic = True
        elif token.startswith("["):
            label, target = token[1:].split("](", 1)
            target = target[:-1]
            _append_styled_run(paragraph, label or target, size=default_size)
            if target and target != label:
                run = _append_styled_run(paragraph, f"（{target}）", size=default_size)
                run.font.color.rgb = RGBColor(89, 89, 89)
        cursor = match.end()
    if cursor < len(text):
        _append_styled_run(paragraph, text[cursor:], size=default_size)


def _append_styled_run(paragraph, text: str, *, size: float | None = None, font: str = "宋体"):
    run = paragraph.add_run(text)
    run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    _set_run_east_asia_font(run, font)
    return run


def _add_code_block(document: DocxDocument, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(6)
    _set_paragraph_shading(paragraph, "F2F2F2")
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    _set_run_east_asia_font(run, "等线")


def _join_soft_lines(lines: list[str]) -> str:
    result = ""
    for value in lines:
        if not result:
            result = value
            continue
        separator = " " if _needs_space(result[-1], value[0]) else ""
        result += separator + value
    return result


def _needs_space(left: str, right: str) -> bool:
    return bool(re.match(r"[A-Za-z0-9%)]", left) and re.match(r"[A-Za-z0-9(]", right))


def _strip_inline_markdown(text: str) -> str:
    text = re.sub(r"!\[([^]]*)]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)]\([^)]+\)", r"\1", text)
    text = re.sub(r"(\*\*|__|`|\*|_)", "", text)
    return text.strip()


def _looks_numeric(value: str) -> bool:
    clean = _strip_inline_markdown(value).strip().replace(",", "").replace("%", "")
    clean = clean.replace("—", "").replace("-", "", 1)
    return bool(clean) and bool(re.fullmatch(r"[()\d.]+", clean))


def _set_east_asia_font(style, font_name: str) -> None:
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _set_run_east_asia_font(run, font_name: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _append_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def _create_numbering_instance(document: DocxDocument, *, start_at: int) -> int:
    """Create a fresh Word numbering instance so each Markdown list restarts."""

    numbering = document.part.numbering_part.element
    base_num_id = str(document.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(
        item
        for item in numbering.findall(qn("w:num"))
        if item.get(qn("w:numId")) == base_num_id
    )
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [
        int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(start_at))
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    value = OxmlElement("w:numId")
    value.set(qn("w:val"), str(num_id))
    num_pr.extend((level, value))
    p_pr.append(num_pr)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E1F2")
        borders.append(element)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def _set_paragraph_left_rule(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "A6A6A6")
    borders.append(left)
    p_pr.append(borders)


__all__ = [
    "MarkdownImageInspection",
    "WordExportError",
    "WordExportResult",
    "export_markdown_report",
    "inspect_markdown_images",
]
