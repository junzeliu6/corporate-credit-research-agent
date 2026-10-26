from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image
import pytest

from credit_research_agent.word_export import WordExportError, export_markdown_report


def _write_png(path: Path, *, size: tuple[int, int] = (1200, 600)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color=(235, 242, 248)).save(path)


def _sample_markdown(image_target: str = "assets/ownership_structure.svg") -> str:
    return f"""# 石家庄以岭药业股份有限公司企业授信研究报告

> 本报告为 AI 生成研究初稿，需人工复核。

## 第一章：企业基本情况

公司**主营业务**包括中成药，并保留 `002603.SZ` 标识。
这是中文软换行，应该与上一行合并。

### 发展历史

- 1992年开始相关业务
- 2001年设立股份公司

1. 核对原始披露
2. 复核重要假设

## 股权结构图

![股权结构图]({image_target})

## 财务分析

| 科目 | FY2024（万元） | FY2025（万元） | 同比 |
| --- | ---: | ---: | ---: |
| 营业收入 | 1,000.00 | 1,200.00 | 20.00% |
| 经营现金流 | 100.00 | 290.75 | 190.75% |

## 待补充资料及尽调问题

- 请补充授信明细及到期安排。
"""


def test_exports_real_word_structures_and_prefers_png(tmp_path: Path) -> None:
    root = tmp_path / "task"
    markdown = root / "credit_research_report.md"
    png = root / "assets" / "ownership_structure.png"
    svg = root / "assets" / "ownership_structure.svg"
    _write_png(png)
    svg.parent.mkdir(parents=True, exist_ok=True)
    svg.write_text("<svg/>", encoding="utf-8")
    original = _sample_markdown()
    markdown.write_text(original, encoding="utf-8")

    result = export_markdown_report(
        markdown,
        allowed_root=root,
        company_name="石家庄以岭药业股份有限公司",
        ticker="002603.SZ",
        report_date="2026-10-26",
        require_ownership_image=True,
    )

    assert result.docx_path == markdown.with_suffix(".docx").resolve()
    assert result.docx_path.is_file()
    assert result.image_count == 1
    assert result.table_count == 1
    assert result.heading_count == 5
    assert markdown.read_text(encoding="utf-8") == original

    document = Document(result.docx_path)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "石家庄以岭药业股份有限公司企业授信研究报告" in text
    assert "本报告由 AI 辅助生成，相关结论需人工复核。" in text
    assert "本报告为 AI 生成研究初稿，需人工复核。" in text
    assert "主营业务" in text
    assert "请补充授信明细及到期安排" in text
    assert any(p.style.name == "Heading 2" and "第一章" in p.text for p in document.paragraphs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == "1,000.00"
    assert len(document.inline_shapes) == 1
    assert "PAGE" in document.sections[0].footer._element.xml
    assert document.styles["Title"]._element.pPr.find(qn("w:pBdr")) is None


def test_explicit_ownership_image_is_inserted_after_heading(tmp_path: Path) -> None:
    root = tmp_path / "task"
    root.mkdir(parents=True)
    markdown = root / "report.md"
    markdown.write_text("# 报告\n\n## 股权结构图\n\n股权说明。\n", encoding="utf-8")
    image = root / "assets" / "ownership_structure.png"
    _write_png(image, size=(600, 1200))

    result = export_markdown_report(
        markdown,
        allowed_root=root,
        ownership_image_path=image,
        require_ownership_image=True,
    )

    assert result.image_count == 1
    assert result.warning_codes == ()
    shape = Document(result.docx_path).inline_shapes[0]
    assert shape.width <= Cm(15.4)
    assert shape.height <= Cm(20.0)


def test_explicit_ownership_image_replaces_markdown_reference(tmp_path: Path) -> None:
    root = tmp_path / "task"
    root.mkdir(parents=True)
    markdown = root / "report.md"
    markdown.write_text(
        "# 报告\n\n## 股权结构图\n\n![股权结构图](assets/ownership_structure.svg)\n",
        encoding="utf-8",
    )
    image = root / "assets" / "ownership_structure.png"
    _write_png(image)

    result = export_markdown_report(
        markdown,
        allowed_root=root,
        ownership_image_path=image,
        require_ownership_image=True,
    )

    assert result.image_count == 1
    assert len(Document(result.docx_path).inline_shapes) == 1


def test_missing_referenced_image_fails_clearly(tmp_path: Path) -> None:
    root = tmp_path / "task"
    root.mkdir()
    markdown = root / "report.md"
    markdown.write_text(_sample_markdown("assets/missing.png"), encoding="utf-8")

    with pytest.raises(WordExportError) as error:
        export_markdown_report(markdown, allowed_root=root)

    assert error.value.code == "referenced_image_missing"
    assert not markdown.with_suffix(".docx").exists()


@pytest.mark.parametrize(
    ("image_target", "expected_code"),
    [
        ("https://example.com/ownership.png", "remote_image_not_allowed"),
        ("../outside.png", "image_outside_allowed_root"),
    ],
)
def test_rejects_untrusted_image_paths(
    tmp_path: Path, image_target: str, expected_code: str
) -> None:
    root = tmp_path / "task"
    root.mkdir()
    markdown = root / "report.md"
    markdown.write_text(_sample_markdown(image_target), encoding="utf-8")
    _write_png(tmp_path / "outside.png")

    with pytest.raises(WordExportError) as error:
        export_markdown_report(markdown, allowed_root=root)

    assert error.value.code == expected_code


def test_rejects_output_outside_task_directory(tmp_path: Path) -> None:
    root = tmp_path / "task"
    root.mkdir()
    markdown = root / "report.md"
    markdown.write_text("# 报告\n\n正文。", encoding="utf-8")

    with pytest.raises(WordExportError) as error:
        export_markdown_report(markdown, tmp_path / "outside.docx", allowed_root=root)

    assert error.value.code == "output_outside_allowed_root"


def test_required_ownership_image_is_not_silently_omitted(tmp_path: Path) -> None:
    root = tmp_path / "task"
    root.mkdir()
    markdown = root / "report.md"
    markdown.write_text("# 报告\n\n## 股权结构图\n\n图表待补充。", encoding="utf-8")

    with pytest.raises(WordExportError) as error:
        export_markdown_report(markdown, allowed_root=root, require_ownership_image=True)

    assert error.value.code == "ownership_image_required"
    assert not markdown.with_suffix(".docx").exists()


def test_separate_ordered_lists_restart_and_identifier_underscores_stay_plain(
    tmp_path: Path,
) -> None:
    root = tmp_path / "task"
    root.mkdir()
    markdown = root / "report.md"
    markdown.write_text(
        "# 报告\n\n1. 第一项\n2. 第二项\n\n## 新章节\n\n"
        "1. 新清单第一项\n2. 新清单第二项\n\n"
        "- fetch_source:source_not_from_search\n",
        encoding="utf-8",
    )

    result = export_markdown_report(markdown, allowed_root=root)
    document = Document(result.docx_path)
    numbered = [p for p in document.paragraphs if p.style.name == "List Number"]
    num_ids = [
        p._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        for p in numbered
    ]

    assert len(numbered) == 4
    assert num_ids[0] == num_ids[1]
    assert num_ids[2] == num_ids[3]
    assert num_ids[0] != num_ids[2]
    identifier = next(
        p for p in document.paragraphs if "fetch_source:source_not_from_search" in p.text
    )
    assert not any(run.italic for run in identifier.runs)
