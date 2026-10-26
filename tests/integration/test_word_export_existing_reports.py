from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document

from credit_research_agent.word_export import export_markdown_report


PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXISTING_REPORTS = (
    PROJECT_ROOT
    / "data"
    / "complete_research_live_v25_retest_20261026"
    / "CN-A-SZSE-002603"
    / "credit_research_report_recovered.md",
    PROJECT_ROOT
    / "data"
    / "streamlit_runs"
    / "v31_mock_e2e_20261026"
    / "CN-A-SZSE-002603"
    / "credit_research_report.md",
)


@pytest.mark.integration
@pytest.mark.parametrize("source_report", EXISTING_REPORTS, ids=("v25-live-recovered", "v31-mock"))
def test_existing_complete_report_exports_with_graph_and_financial_tables(
    source_report: Path, tmp_path: Path
) -> None:
    if not source_report.is_file():
        pytest.skip(f"本机既有报告不存在：{source_report}")

    task_dir = tmp_path / source_report.parent.name
    assets_dir = task_dir / "assets"
    assets_dir.mkdir(parents=True)
    report = task_dir / source_report.name
    shutil.copy2(source_report, report)
    for suffix in (".png", ".svg"):
        source_image = source_report.parent / "assets" / f"ownership_structure{suffix}"
        if source_image.is_file():
            shutil.copy2(source_image, assets_dir / source_image.name)

    png = assets_dir / "ownership_structure.png"
    assert png.is_file(), "现有完整报告应保留可嵌入Word的股权PNG"

    result = export_markdown_report(
        report,
        allowed_root=task_dir,
        ownership_image_path=png,
        company_name="石家庄以岭药业股份有限公司",
        ticker="002603.SZ",
        report_date="2026-10-26",
        require_ownership_image=True,
    )

    document = Document(result.docx_path)
    paragraphs = "\n".join(item.text for item in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "第一章：企业基本情况" in paragraphs
    assert "本报告由 AI 辅助生成，相关结论需人工复核。" in paragraphs
    assert result.image_count == 1
    assert len(document.inline_shapes) == 1
    assert result.table_count >= 2
    assert "营业收入" in table_text
    assert "20.24%" in table_text
